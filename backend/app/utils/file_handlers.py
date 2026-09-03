import os
import io
import csv
import json
import PyPDF2
import docx
from pptx import Presentation
from openpyxl import load_workbook
from bs4 import BeautifulSoup
import markdown
from typing import Dict, Any
from fastapi import HTTPException


def extract_text_from_file(content: bytes, filename: str) -> str:
    
    
    # Get file extension
    ext = filename.split('.')[-1].lower()
    
    try:
        if ext == 'pdf':
            return extract_pdf(content)
        elif ext == 'docx':
            return extract_docx(content)
        elif ext == 'txt':
            return extract_txt(content)
        elif ext == 'html' or ext == 'htm':
            return extract_html(content)
        elif ext == 'csv':
            return extract_csv(content)
        elif ext == 'json':
            return extract_json(content)
        elif ext == 'md':
            return extract_markdown(content)
        elif ext == 'pptx':
            return extract_pptx(content)
        elif ext == 'xlsx':
            return extract_xlsx(content)
        else:
            # Try reading as text for unknown formats
            return extract_txt(content)
    
    except Exception as e:
        raise HTTPException(
            status_code=400,
            detail=f"Failed to extract text from {filename}: {str(e)}"
        )


def extract_pdf(content: bytes) -> str:
    
    pdf_file = io.BytesIO(content)
    reader = PyPDF2.PdfReader(pdf_file)
    text = ""
    for page in reader.pages:
        text += page.extract_text() + "\n"
    return text.strip()


def extract_docx(content: bytes) -> str:
    
    doc_file = io.BytesIO(content)
    doc = docx.Document(doc_file)
    text = ""
    for para in doc.paragraphs:
        text += para.text + "\n"
    return text.strip()


def extract_txt(content: bytes) -> str:
    
    try:
        return content.decode('utf-8')
    except:
        try:
            return content.decode('latin-1')
        except:
            return content.decode('utf-8', errors='ignore')


def extract_html(content: bytes) -> str:
   
    html = content.decode('utf-8', errors='ignore')
    soup = BeautifulSoup(html, 'html.parser')
    # Remove script and style tags
    for script in soup(["script", "style"]):
        script.decompose()
    text = soup.get_text()
    # Clean whitespace
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    return "\n".join(lines)


def extract_csv(content: bytes) -> str:
    
    text = content.decode('utf-8', errors='ignore')
    lines = text.splitlines()
    reader = csv.reader(lines)
    result = []
    for row in reader:
        result.append(" | ".join(row))
    return "\n".join(result)


def extract_json(content: bytes) -> str:
    
    text = content.decode('utf-8', errors='ignore')
    data = json.loads(text)
    return json.dumps(data, indent=2)


def extract_markdown(content: bytes) -> str:
    
    text = content.decode('utf-8', errors='ignore')
    # Convert markdown to HTML, then extract text
    html = markdown.markdown(text)
    soup = BeautifulSoup(html, 'html.parser')
    return soup.get_text().strip()


def extract_pptx(content: bytes) -> str:
   
    pptx_file = io.BytesIO(content)
    prs = Presentation(pptx_file)
    text = ""
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text += shape.text + "\n"
    return text.strip()


def extract_xlsx(content: bytes) -> str:
    
    xlsx_file = io.BytesIO(content)
    wb = load_workbook(xlsx_file, data_only=True)
    text = ""
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        text += f"\n--- Sheet: {sheet} ---\n"
        for row in ws.iter_rows(values=True):
            row_text = " | ".join([str(cell) if cell is not None else "" for cell in row])
            if row_text.strip():
                text += row_text + "\n"
    return text.strip()


def get_file_metadata(filename: str, content: bytes) -> Dict[str, Any]:
    
    ext = filename.split('.')[-1].lower() if '.' in filename else 'unknown'
    
    return {
        "filename": filename,
        "extension": ext,
        "size_bytes": len(content),
        "size_kb": round(len(content) / 1024, 2),
        "type": get_file_type(ext)
    }


def get_file_type(ext: str) -> str:
    type_map = {
        'pdf': 'PDF Document',
        'docx': 'Word Document',
        'txt': 'Text File',
        'html': 'HTML File',
        'htm': 'HTML File',
        'csv': 'CSV File',
        'json': 'JSON File',
        'md': 'Markdown File',
        'pptx': 'PowerPoint File',
        'xlsx': 'Excel File',
    }
    return type_map.get(ext, 'Unknown')